import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import io
import pandas as pd
from typing import Dict, List, Any

class ExcelService:
    def create_template(self, table_name: str, columns: List[Dict[str, Any]]) -> bytes:
        wb = openpyxl.Workbook()
        ws_data = wb.active
        ws_data.title = "Data_Upload_Template"
        
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        required_header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        type_subfont = Font(name="Calibri", size=9, italic=True, color="64748B")
        
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        
        col_idx = 1
        for col in columns:
            col_name = col["column"]
            data_type = col.get("type", "varchar").upper()
            nullable = col.get("nullable", "YES") == "YES"
            
            if col.get("identity", False):
                continue
                
            display_name = f"{col_name} *" if not nullable else col_name
            cell = ws_data.cell(row=1, column=col_idx, value=display_name)
            cell.font = header_font
            cell.fill = required_header_fill if not nullable else header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            
            type_label = f"{data_type}"
            if col.get("length") and col["length"] > 0:
                type_label += f"({col['length']})"
            cell_type = ws_data.cell(row=2, column=col_idx, value=type_label)
            cell_type.font = type_subfont
            cell_type.alignment = Alignment(horizontal="center", vertical="center")
            cell_type.border = thin_border
            
            col_idx += 1
            
        ws_data.row_dimensions[1].height = 28
        ws_data.row_dimensions[2].height = 20
        
        # Instructions Sheet
        ws_inst = wb.create_sheet(title="Instructions_Guide")
        ws_inst.cell(row=1, column=1, value=f"DK's SQL Buddy - Excel Import Guide for [{table_name}]").font = Font(size=14, bold=True, color="1E3A8A")
        
        instructions = [
            ("1. Mandatory Columns", "Columns marked with asterisk (*) are required and cannot be left empty."),
            ("2. Data Format Rules", "Dates must be in YYYY-MM-DD format. Numbers must contain numeric values only."),
            ("3. String Length Limits", "Text entries must not exceed the maximum character length specified in row 2."),
            ("4. Header Integrity", "Do NOT rename, reorder, or delete column header names in Row 1."),
            ("5. Auto-ID Fields", "System identity/Primary Key fields are handled automatically during import.")
        ]
        
        row_num = 3
        for title, desc in instructions:
            c1 = ws_inst.cell(row=row_num, column=1, value=title)
            c1.font = Font(bold=True, size=11, color="1E293B")
            c2 = ws_inst.cell(row=row_num, column=2, value=desc)
            c2.font = Font(size=11, color="334155")
            row_num += 1
            
        for ws in [ws_data, ws_inst]:
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 15)
                
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def export_table_data(self, table_name: str, columns: List[Dict[str, Any]], data_rows: List[Dict[str, Any]]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Exported_Table_Data"
        
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        
        col_names = [c["column"] for c in columns]
        
        # Write Headers
        for col_idx, col_name in enumerate(col_names, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        ws.row_dimensions[1].height = 28
        
        # Write Rows
        for row_idx, record in enumerate(data_rows, start=2):
            for col_idx, col_name in enumerate(col_names, start=1):
                val = record.get(col_name)
                ws.cell(row=row_idx, column=col_idx, value=str(val) if val is not None else "")
                
        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 40)
            
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def parse_excel(self, file_contents: bytes) -> List[Dict[str, Any]]:
        wb = openpyxl.load_workbook(io.BytesIO(file_contents), data_only=True)
        sheet_name = wb.sheetnames[0]
        for name in wb.sheetnames:
            if "data" in name.lower() or "template" in name.lower() or "export" in name.lower():
                sheet_name = name
                break
                
        ws = wb[sheet_name]
        
        headers = []
        for cell in ws[1]:
            val = str(cell.value or "").strip()
            if val.endswith("*"):
                val = val[:-1].strip()
            if val:
                headers.append(val)
                
        if not headers:
            raise ValueError("No header columns found in Excel file.")
            
        rows = []
        start_row = 2
        r2_vals = [str(c.value or "").strip() for c in ws[2]]
        if any(t in "".join(r2_vals).upper() for t in ["VARCHAR", "NVARCHAR", "INT", "DATE", "DECIMAL", "CHAR"]):
            start_row = 3
            
        for row_idx in range(start_row, ws.max_row + 1):
            row_data = {}
            has_data = False
            for col_idx, header in enumerate(headers, start=1):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None and str(val).strip() != "":
                    has_data = True
                row_data[header] = val
                
            if has_data:
                rows.append(row_data)
                
        return rows

excel_service = ExcelService()
